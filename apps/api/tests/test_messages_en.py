"""L'API repond dans la langue de la page qui l'appelle.

Ce que ces tests tiennent :

  - chaque message que le code peut renvoyer a sa traduction anglaise —
    un message ajoute demain et oublie dans `common/messages_en.py` fait
    echouer ce fichier, au lieu de sortir en francais sur une page anglaise ;
  - une requete en anglais recoit ses refus en anglais, ceux du produit
    comme ceux de Django et de DRF ; une requete en francais n'est pas
    touchee ;
  - seuls les messages sont traduits, jamais les donnees ;
  - le fil d'activite et les comptes donnent leurs faits pour que la page
    compose la phrase dans sa langue.
"""

import ast
import json
import re
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook

from apps.common.langue_api import traduire_reponse
from apps.common.messages_en import EXACTS, MODELES, _propre, traduire
from apps.finance import libelles

RACINE = Path(__file__).resolve().parent.parent / "apps"
PASSWORD = "motdepasse-solide"

#: Visibles seulement d'un developpeur ou dans l'administration (francaise),
#: ou donnees et non messages.
INTERNES = {
    "Un superutilisateur doit avoir is_staff et is_superuser.",
    "Aucun tenant dans le contexte courant. Enveloppez l'appel dans tenant_context(tenant_id).",
    "Cet element appartient a un autre salon.",
    "Une seule politique plateforme est autorisée.",
    (
        "Complétez et validez tous les paramètres avant activation. Cette version accepte "
        "une remise de 10 % par échéance."
    ),
    (
        "Configuration incomplète : taux (10 % minimum), durée, plafond monétaire et "
        "politiques validées sont obligatoires."
    ),
    "Un intervalle doit finir apres son debut.",
    "Longueur, couleur, quantité conseillée…",
    "Ce moyen ne peut pas être actif : %(manques)s.",
    "Rappel inconnu : Anna",
    "Geste inconnu : Anna",
    "Courrier inconnu : Anna",
    "Contexte tenant deja actif (Anna), impossible de basculer vers Anna sans sortir du "
    "bloc courant.",
    # Le detail d'une ligne du fil d'activite : un montant, pas un message.
    "Anna Anna",
}

#: Les fichiers dont une cle « detail » n'est pas un message d'API.
HORS_API = {"excel.py", "admin.py"}
APPELS = (
    "Error", "Exception", "Denied", "NotFound", "Invalid", "Refus", "Conflict",
    "Unavailable", "Throttled", "ParseError", "Inconnu", "Refuse", "Interdit",
)


def _exemple(noeud) -> str | None:
    """La phrase telle qu'elle sortirait, avec des valeurs d'exemple : un
    nom pour chaque variable, une date pour une heure formatee."""
    if isinstance(noeud, ast.Constant) and isinstance(noeud.value, str):
        return noeud.value
    if isinstance(noeud, ast.JoinedStr):
        morceaux = []
        for v in noeud.values:
            if isinstance(v, ast.Constant):
                morceaux.append(v.value)
            elif v.format_spec is not None and "à" in ast.unparse(v.format_spec):
                morceaux.append("24/09 à 14:30")
            else:
                morceaux.append("Anna")
        return "".join(morceaux)
    if isinstance(noeud, ast.BinOp) and isinstance(noeud.op, ast.Add):
        a, b = _exemple(noeud.left), _exemple(noeud.right)
        # « Indiquez ce que vous apportez : » + ", ".join(...) : la partie
        # calculee devient une valeur d'exemple.
        if a is not None and b is None and isinstance(noeud.right, ast.Call):
            b = "Anna"
        if a is not None and b is not None:
            return a + b
    return None


def _phrases(noeud):
    exemple = _exemple(noeud)
    if exemple is not None:
        yield exemple
        return
    if isinstance(noeud, ast.BinOp) and isinstance(noeud.op, ast.Add):
        # « debut » + (« a » if x else « b ») : les deux phrases possibles.
        gauche = _exemple(noeud.left)
        if gauche is not None and isinstance(noeud.right, ast.IfExp):
            for branche in (noeud.right.body, noeud.right.orelse):
                suite = _exemple(branche)
                if suite is not None:
                    yield gauche + suite
            return
    for enfant in ast.iter_child_nodes(noeud):
        yield from _phrases(enfant)


def _messages_du_code() -> set[str]:
    """Chaque message que l'API peut renvoyer : `detail`, exceptions,
    erreurs de validation — tels qu'ils sortiraient."""
    trouves = set()
    for fichier in RACINE.rglob("*.py"):
        if {"migrations", "tests", "management"} & set(fichier.parts) or fichier.name in HORS_API:
            continue
        arbre = ast.parse(fichier.read_text(encoding="utf-8"))
        for noeud in ast.walk(arbre):
            valeurs = []
            if isinstance(noeud, ast.Call):
                f = noeud.func
                nom = f.attr if isinstance(f, ast.Attribute) else getattr(f, "id", "")
                if nom.endswith(APPELS):
                    valeurs = [*noeud.args, *(k.value for k in noeud.keywords)]
            elif isinstance(noeud, ast.Dict):
                valeurs = [
                    v for k, v in zip(noeud.keys, noeud.values, strict=True)
                    if isinstance(k, ast.Constant) and k.value in ("detail", "message")
                ]
            elif isinstance(noeud, ast.Assign):
                if any(
                    isinstance(c, ast.Name) and c.id in ("default_detail", "message", "detail")
                    for c in noeud.targets
                ):
                    valeurs = [noeud.value]
            for valeur in valeurs:
                for phrase in _phrases(valeur):
                    phrase = " ".join(phrase.split())
                    if len(phrase) > 3 and " " in phrase and not phrase.startswith("/"):
                        trouves.add(phrase)
    return trouves


# ---------------------------------------------------------------------------
# Le dictionnaire
# ---------------------------------------------------------------------------


def test_chaque_message_du_code_a_sa_traduction():
    oublies = sorted(
        m for m in _messages_du_code()
        if m not in INTERNES and traduire(m) == m
    )
    assert not oublies, "Messages sans traduction anglaise :\n" + "\n".join(oublies)


def test_le_dictionnaire_ne_garde_pas_de_message_disparu():
    presents = {_propre(m) for m in _messages_du_code()}
    disparus = [fr for fr in EXACTS if _propre(fr) not in presents]
    assert not disparus, "Traductions sans message dans le code :\n" + "\n".join(disparus)


def test_les_modeles_gardent_leurs_variables():
    for francais, anglais in MODELES:
        variables = re.findall(r"\{(\w+)\}", francais)
        assert sorted(variables) == sorted(re.findall(r"\{(\w+)\}", anglais)), francais


@pytest.mark.parametrize(
    "francais, anglais",
    [
        ("Mot de passe incorrect.", "Incorrect password."),
        # L'apostrophe typographique ne fait pas manquer la traduction.
        ("Ce paiement n’existe pas.", "This payment does not exist."),
        ("« Huile d'argan » : il n'en reste que 2.", "“Huile d'argan”: only 2 left."),
        (
            "Ce rendez-vous n'est pas confirmé : il est « annulée ».",
            "This appointment is not confirmed: it is “cancelled”.",
        ),
        (
            "Ce code est celui de Anna Bika, pas de Grâce M.",
            "This code belongs to Anna Bika, not to Grâce M.",
        ),
        ("Un texte que personne n'a prévu.", "Un texte que personne n'a prévu."),
    ],
)
def test_traduire(francais, anglais):
    assert traduire(francais) == anglais


def test_seuls_les_messages_sont_traduits():
    reponse = {
        "detail": "Ce créneau est déjà passé.",
        "code": "slot_past",
        # Une donnee jointe au refus : elle reste ce qu'elle est.
        "booking": {"customer_note": "Date invalide."},
    }
    traduite = traduire_reponse(reponse, erreur=True)
    assert traduite["detail"] == "This time slot has already passed."
    assert traduite["code"] == "slot_past"
    assert traduite["booking"] == {"customer_note": "Date invalide."}


def test_une_erreur_de_validation_brute_est_traduite_champ_par_champ():
    traduite = traduire_reponse({"quantity": ["Quantité invalide."]}, erreur=True)
    assert traduite == {"quantity": ["Invalid quantity."]}


# ---------------------------------------------------------------------------
# De bout en bout
# ---------------------------------------------------------------------------


@pytest.mark.django_db
def test_un_refus_part_en_anglais_pour_une_page_anglaise(api_client, salon_a):
    reponse = api_client.post(
        "/api/v1/auth/login",
        {"email": salon_a.owner.email, "password": "mauvais"},
        format="json",
        HTTP_ACCEPT_LANGUAGE="en",
    )

    assert reponse.status_code >= 400
    assert reponse.json()["detail"] == "Incorrect email address or password."


@pytest.mark.django_db
def test_un_refus_reste_en_francais_pour_une_page_francaise(api_client, salon_a):
    reponse = api_client.post(
        "/api/v1/auth/login",
        {"email": salon_a.owner.email, "password": "mauvais"},
        format="json",
        HTTP_ACCEPT_LANGUAGE="fr",
    )

    assert reponse.json()["detail"] == "Adresse e-mail ou mot de passe incorrect."


@pytest.mark.django_db
def test_les_messages_de_django_et_drf_suivent_la_page(api_client):
    anglais = api_client.post(
        "/api/v1/account/password/reset", {}, format="json", HTTP_ACCEPT_LANGUAGE="en"
    )
    francais = api_client.post(
        "/api/v1/account/password/reset", {}, format="json", HTTP_ACCEPT_LANGUAGE="fr"
    )

    assert "This field is required." in json.dumps(anglais.json())
    assert "Ce champ est obligatoire." in json.dumps(francais.json(), ensure_ascii=False)


@pytest.mark.django_db
def test_une_reponse_reussie_garde_ses_donnees(api_client, salon_a):
    assert api_client.login(email=salon_a.owner.email, password=PASSWORD)

    reponse = api_client.get(
        "/api/v1/auth/session", HTTP_ACCEPT_LANGUAGE="en", HTTP_X_TENANT_ID=str(salon_a.tenant.id)
    )

    assert reponse.status_code == 200
    assert reponse.json()["email"] == salon_a.owner.email


# ---------------------------------------------------------------------------
# Comptes du salon
# ---------------------------------------------------------------------------


class _Ecriture:
    def __init__(self, label, source="booking_deposit", category="service", method="cash"):
        self.label, self.source, self.category, self.method = label, source, category, method
        self.category_label = "Prestation"

    def get_method_display(self):
        return "Espèces"


def test_les_libelles_automatiques_se_disent_en_anglais():
    assert libelles.libelle(_Ecriture("Acompte — Box braids"), "en") == "Deposit — Box braids"
    assert libelles.libelle(_Ecriture("Vente au comptoir", "booking_items"), "en") == "Counter sale"
    assert (
        libelles.libelle(_Ecriture("Abonnement Beauty Salon — 10/2026", "subscription"), "en")
        == "Beauty Salon subscription — 10/2026"
    )
    # Une ecriture saisie a la main appartient au salon.
    assert libelles.libelle(_Ecriture("Acompte — loyer", "manual"), "en") == "Acompte — loyer"
    assert libelles.libelle(_Ecriture("Acompte — Box braids"), "fr") == "Acompte — Box braids"
    assert libelles.poste(_Ecriture("x"), "en") == "Service"
    assert libelles.moyen(_Ecriture("x"), "en") == "Cash"


@pytest.mark.django_db
def test_l_export_excel_suit_la_langue_de_la_page(api_client, salon_a):
    assert api_client.login(email=salon_a.owner.email, password=PASSWORD)

    reponse = api_client.get(
        "/api/v1/transactions/export/",
        HTTP_ACCEPT_LANGUAGE="en",
        HTTP_X_TENANT_ID=str(salon_a.tenant.id),
    )

    assert reponse.status_code == 200
    classeur = load_workbook(BytesIO(reponse.content))
    assert classeur.sheetnames == ["Summary", "Income", "Expenses"]
    assert classeur["Income"]["B1"].value == "Label"
