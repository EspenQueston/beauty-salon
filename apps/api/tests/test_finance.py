"""Recettes et depenses.

C'est la table la plus sensible du produit apres les fiches clientes : elle
porte le chiffre d'affaires, les charges et la marge. Ces tests verrouillent
trois choses - qui peut la lire, ce qu'elle refuse d'enregistrer, et le fait
qu'une prestation honoree ne soit jamais comptee deux fois.
"""

from contextlib import contextmanager
from datetime import date, datetime, timedelta
from unittest import mock
from zoneinfo import ZoneInfo
from decimal import Decimal

import pytest
from django.utils import timezone

from apps.accounts.models import Membership
from apps.finance.models import Transaction
from apps.finance.services import aujourdhui_du_salon
from apps.scheduling.models import Booking
from conftest import as_tenant, salon_host
from tests.factories import MembershipFactory, UserFactory

PASSWORD = "motdepasse-solide"
HOST = {"Host": salon_host("blondrose")}


def login(client, user):
    assert client.login(email=user.email, password=PASSWORD)
    return client


def aujourdhui(salon) -> date:
    """« Aujourd.hui » du point de vue du salon, pas du poste qui lance les tests.

    Les tests fabriquaient leurs lignes avec `date.today()`, c.est-a-dire la
    date de la machine. Ils passaient sur un poste regle sur le meme fuseau
    que le salon de test - Africa/Brazzaville - et echouaient ailleurs, ou
    bien masquaient le bogue inverse : une fenetre calculee sur une autre
    horloge que les lignes qu.elle filtre.
    """
    return aujourdhui_du_salon(salon.tenant)


def entry(salon, **overrides):
    payload = {
        "kind": Transaction.Kind.EXPENSE,
        "category": Transaction.ExpenseCategory.SUPPLIES,
        "label": "Mèches kanekalon",
        "amount": "12000.00",
        "occurred_on": aujourdhui(salon).isoformat(),
        "method": Transaction.Method.CASH,
    }
    payload.update(overrides)
    return payload


# ---------------------------------------------------------------------------
# Saisie
# ---------------------------------------------------------------------------


@pytest.mark.django_db
def test_an_owner_records_an_expense(api_client, salon_a):
    login(api_client, salon_a.owner)

    response = api_client.post("/api/v1/transactions/", entry(salon_a), format="json")

    assert response.status_code == 201, response.data
    assert response.data["category_label"] == "Fournitures et produits"


@pytest.mark.django_db
def test_a_category_from_the_wrong_side_is_refused(api_client, salon_a):
    """« Loyer » n'est pas une recette, « pourboire » n'est pas une charge."""
    login(api_client, salon_a.owner)

    response = api_client.post(
        "/api/v1/transactions/",
        entry(salon_a, kind=Transaction.Kind.INCOME, category=Transaction.ExpenseCategory.RENT),
        format="json",
    )

    assert response.status_code == 400
    assert "category" in str(response.data)


@pytest.mark.django_db
def test_a_negative_amount_is_refused(api_client, salon_a):
    """Le sens est porte par `kind`, jamais par le signe du montant.

    Stocker les depenses en negatif est tentant et se paie cher : un signe
    perdu transforme une charge en recette, sans rien pour le detecter.
    """
    login(api_client, salon_a.owner)

    response = api_client.post(
        "/api/v1/transactions/", entry(salon_a, amount="-500.00"), format="json"
    )

    assert response.status_code == 400


@pytest.mark.django_db
def test_a_future_date_is_refused(api_client, salon_a):
    """Une depense a venir est une prevision, pas un mouvement."""
    login(api_client, salon_a.owner)

    response = api_client.post(
        "/api/v1/transactions/",
        entry(
            salon_a,
            occurred_on=(aujourdhui(salon_a) + timedelta(days=3)).isoformat(),
        ),
        format="json",
    )

    assert response.status_code == 400
    assert "occurred_on" in str(response.data)


# ---------------------------------------------------------------------------
# Qui a le droit de regarder
# ---------------------------------------------------------------------------


@pytest.mark.django_db
@pytest.mark.parametrize(
    "role", [Membership.Role.STAFF, Membership.Role.RECEPTIONIST]
)
def test_the_team_cannot_read_the_accounts(api_client, salon_a, role):
    """Un prestataire voit son agenda, pas la marge du mois ni le loyer.

    La reception est exclue elle aussi : elle encaisse, elle ne tient pas la
    comptabilite.
    """
    member = UserFactory()
    MembershipFactory(tenant=salon_a.tenant, user=member, role=role)
    login(api_client, member)

    listing = api_client.get("/api/v1/transactions/")
    summary = api_client.get("/api/v1/transactions/summary/")

    assert listing.status_code == 403
    assert summary.status_code == 403


@pytest.mark.django_db
def test_one_salon_never_sees_another_salons_accounts(api_client, salon_a, salon_b):
    with as_tenant(salon_b.tenant):
        Transaction.objects.create(
            tenant=salon_b.tenant,
            kind=Transaction.Kind.EXPENSE,
            category=Transaction.ExpenseCategory.RENT,
            label="Loyer du salon B",
            amount=Decimal("300000"),
            occurred_on=aujourdhui(salon_a),
        )

    login(api_client, salon_a.owner)
    response = api_client.get("/api/v1/transactions/")

    assert response.status_code == 200
    assert response.data["results"] == []


# ---------------------------------------------------------------------------
# Synthese
# ---------------------------------------------------------------------------


@pytest.mark.django_db
def test_the_summary_adds_up(api_client, salon_a):
    with as_tenant(salon_a.tenant):
        for amount, kind, category in [
            ("100000", Transaction.Kind.INCOME, Transaction.IncomeCategory.SERVICE),
            ("20000", Transaction.Kind.INCOME, Transaction.IncomeCategory.PRODUCT),
            ("30000", Transaction.Kind.EXPENSE, Transaction.ExpenseCategory.SUPPLIES),
            ("10000", Transaction.Kind.EXPENSE, Transaction.ExpenseCategory.RENT),
        ]:
            Transaction.objects.create(
                tenant=salon_a.tenant,
                kind=kind,
                category=category,
                label="ligne",
                amount=Decimal(amount),
                occurred_on=aujourdhui(salon_a),
            )

    login(api_client, salon_a.owner)
    response = api_client.get("/api/v1/transactions/summary/")

    assert response.status_code == 200
    totals = response.data["totals"]
    assert Decimal(totals["income"]) == Decimal("120000")
    assert Decimal(totals["expense"]) == Decimal("40000")
    assert Decimal(totals["net"]) == Decimal("80000")
    # 80 000 / 120 000 = 66,7 %
    assert totals["margin"] == pytest.approx(66.7, abs=0.1)


@pytest.mark.django_db
def test_the_margin_is_unknown_rather_than_zero_without_income(api_client, salon_a):
    """Zero divise par zero vaut « pas de reponse », pas « 0 % »."""
    with as_tenant(salon_a.tenant):
        Transaction.objects.create(
            tenant=salon_a.tenant,
            kind=Transaction.Kind.EXPENSE,
            category=Transaction.ExpenseCategory.RENT,
            label="Loyer",
            amount=Decimal("50000"),
            occurred_on=aujourdhui(salon_a),
        )

    login(api_client, salon_a.owner)
    response = api_client.get("/api/v1/transactions/summary/")

    assert response.data["totals"]["margin"] is None


@pytest.mark.django_db
def test_the_monthly_series_has_no_gaps(api_client, salon_a):
    """Une courbe qui saute de mars a juin laisse croire a une activite
    continue entre les deux."""
    login(api_client, salon_a.owner)

    response = api_client.get("/api/v1/transactions/summary/")

    months = [row["month"] for row in response.data["monthly"]]
    assert len(months) == 12
    assert months == sorted(months), "les mois doivent être dans l'ordre"


@pytest.mark.django_db
def test_the_donut_never_carries_more_than_four_slices(api_client, salon_a):
    """Au-dela de trois teintes, deux parts deviennent indistinguables.

    Le reste est regroupe sous « Autres » — le total, lui, reste exact.
    """
    with as_tenant(salon_a.tenant):
        for category in Transaction.IncomeCategory.values:
            Transaction.objects.create(
                tenant=salon_a.tenant,
                kind=Transaction.Kind.INCOME,
                category=category,
                label=category,
                amount=Decimal("1000"),
                occurred_on=aujourdhui(salon_a),
            )

    login(api_client, salon_a.owner)
    response = api_client.get("/api/v1/transactions/summary/")

    slices = response.data["income_slices"]
    assert len(slices) <= 4
    total = sum(Decimal(row["total"]) for row in slices)
    assert total == Decimal(response.data["totals"]["income"])


# ---------------------------------------------------------------------------
# Recette automatique
# ---------------------------------------------------------------------------


@pytest.mark.django_db
def test_finishing_an_appointment_records_its_income(api_client, salon_a):
    starts_at = timezone.now() - timedelta(days=1)
    with as_tenant(salon_a.tenant):
        booking = Booking.objects.create(
            tenant=salon_a.tenant,
            customer=salon_a.customer,
            staff_member=salon_a.staff,
            service=salon_a.service,
            starts_at=starts_at,
            ends_at=starts_at + timedelta(hours=2),
            status=Booking.Status.CHECKED_IN,
            service_name="Box braids",
            total_amount=Decimal("45000"),
        )

    login(api_client, salon_a.owner)
    response = api_client.post(
        f"/api/v1/bookings/{booking.id}/status/",
        {"status": Booking.Status.COMPLETED},
        format="json",
    )

    assert response.status_code == 200
    with as_tenant(salon_a.tenant):
        recorded = Transaction.objects.get(booking=booking)
    assert recorded.amount == Decimal("45000")
    assert recorded.kind == Transaction.Kind.INCOME
    assert recorded.label == "Box braids"


@pytest.mark.django_db
def test_an_appointment_is_never_counted_twice(api_client, salon_a):
    """Repasser par « terminee » apres une correction ne double pas la caisse."""
    starts_at = timezone.now() - timedelta(days=1)
    with as_tenant(salon_a.tenant):
        booking = Booking.objects.create(
            tenant=salon_a.tenant,
            customer=salon_a.customer,
            staff_member=salon_a.staff,
            service=salon_a.service,
            starts_at=starts_at,
            ends_at=starts_at + timedelta(hours=2),
            status=Booking.Status.CHECKED_IN,
            service_name="Box braids",
            total_amount=Decimal("45000"),
        )

    login(api_client, salon_a.owner)
    for _ in range(3):
        api_client.post(
            f"/api/v1/bookings/{booking.id}/status/",
            {"status": Booking.Status.COMPLETED},
            format="json",
        )

    with as_tenant(salon_a.tenant):
        assert Transaction.objects.filter(booking=booking).count() == 1


# ---------------------------------------------------------------------------
# Export
# ---------------------------------------------------------------------------


@pytest.mark.django_db
def test_the_export_is_one_workbook_with_both_sheets(api_client, salon_a):
    from io import BytesIO

    from openpyxl import load_workbook

    with as_tenant(salon_a.tenant):
        Transaction.objects.create(
            tenant=salon_a.tenant,
            kind=Transaction.Kind.INCOME,
            category=Transaction.IncomeCategory.SERVICE,
            label="Box braids",
            amount=Decimal("45000"),
            occurred_on=aujourdhui(salon_a),
        )
        Transaction.objects.create(
            tenant=salon_a.tenant,
            kind=Transaction.Kind.EXPENSE,
            category=Transaction.ExpenseCategory.RENT,
            label="Loyer",
            amount=Decimal("30000"),
            occurred_on=aujourdhui(salon_a),
        )

    login(api_client, salon_a.owner)
    response = api_client.get("/api/v1/transactions/export/")

    assert response.status_code == 200
    assert "spreadsheetml" in response["Content-Type"]
    assert ".xlsx" in response["Content-Disposition"]

    workbook = load_workbook(BytesIO(response.content))
    assert workbook.sheetnames == ["Synthèse", "Recettes", "Dépenses"]

    # Chaque feuille porte bien ses propres lignes, pas celles de l'autre.
    assert workbook["Recettes"].cell(row=2, column=2).value == "Box braids"
    assert workbook["Dépenses"].cell(row=2, column=2).value == "Loyer"
    # Le montant reste un nombre : c'est ce qui permet de trier et
    # d'additionner dans le tableur.
    assert workbook["Recettes"].cell(row=2, column=6).value == 45000


@pytest.mark.django_db
def test_the_export_is_closed_to_the_team(api_client, salon_a):
    member = UserFactory()
    MembershipFactory(
        tenant=salon_a.tenant, user=member, role=Membership.Role.RECEPTIONIST
    )
    login(api_client, member)

    assert api_client.get("/api/v1/transactions/export/").status_code == 403


# ---------------------------------------------------------------------------
# Le fuseau du salon
# ---------------------------------------------------------------------------
#
# Ces tests verrouillent un bogue invisible en developpement et quotidien en
# production : la journee comptable etait celle du serveur, pas celle du
# salon.
#
# A Shanghai il est deja 1 h du matin le 16 quand UTC en est encore au 15. Une
# prestation terminee le 16 tombait alors hors de la fenetre « jusqu'a
# aujourd'hui », et l'ecran des comptes affichait un total ampute - sans rien
# signaler, puisque de son point de vue la ligne n'existait pas.
#
# ---------------------------------------------------------------------------
# Pourquoi l'horloge est figee dans le futur
# ---------------------------------------------------------------------------
#
# `date.today()` lit l'horloge du systeme, que `mock.patch` sur
# `timezone.now` ne touche pas. Un instant fige proche d'aujourd'hui donnerait
# donc la meme date des deux cotes, et le test passerait aussi bien avec le
# bogue que sans - ce qui ne verrouille rien.
#
# L'instant choisi est donc nettement posterieur a toute date de lancement
# plausible : le salon y est en 2027 quand l'horloge du systeme est encore en
# 2026. La session est ouverte *sous* cette horloge, sans quoi elle serait
# consideree comme expiree et la requete finirait en 403.


def horloge_de_shanghai(salon):
    """Fige l'horloge a 17 h UTC le 31 decembre - 1 h du matin le 1er janvier
    a Shanghai - et regle le salon sur ce fuseau."""
    salon.tenant.timezone = "Asia/Shanghai"
    salon.tenant.save(update_fields=["timezone"])
    return mock.patch(
        "django.utils.timezone.now",
        return_value=datetime(2026, 12, 31, 17, 0, tzinfo=ZoneInfo("UTC")),
    )


@pytest.mark.django_db
def test_the_income_lands_on_the_salon_day_not_the_utc_day(salon_a):
    """Une prestation du matin a Shanghai ne compte pas pour la veille.

    07 h a Shanghai, c'est 23 h UTC la veille. La recette doit porter la date
    que le salon lit sur son mur, sinon il cherche dans ses comptes du mardi
    l'argent qu'il a gagne le mercredi.

    Le rendez-vous est **relu depuis la base** avant d'etre constate, comme
    le fait la vue : c'est seulement a ce moment que Django rend l'instant en
    UTC, et donc seulement a ce moment que le bogue apparaissait.
    """
    from apps.finance.services import record_booking_income

    salon_a.tenant.timezone = "Asia/Shanghai"
    salon_a.tenant.save(update_fields=["timezone"])

    matin = datetime(2026, 9, 16, 7, 0, tzinfo=ZoneInfo("Asia/Shanghai"))
    assert matin.astimezone(ZoneInfo("UTC")).date() == date(2026, 9, 15)

    with as_tenant(salon_a.tenant):
        booking = Booking.objects.create(
            tenant=salon_a.tenant,
            customer=salon_a.customer,
            staff_member=salon_a.staff,
            service=salon_a.service,
            starts_at=matin,
            ends_at=matin + timedelta(hours=2),
            status=Booking.Status.COMPLETED,
            service_name="Box braids",
            total_amount=Decimal("25000"),
        )
        booking.refresh_from_db()
        recette = record_booking_income(booking)

    assert recette is not None
    assert recette.occurred_on == date(2026, 9, 16)


@pytest.mark.django_db
def test_todays_income_stays_visible_whatever_the_server_clock(api_client, salon_a):
    """Le symptome tel qu'il a ete signale : deux lignes en base, une a l'ecran."""
    with horloge_de_shanghai(salon_a):
        with as_tenant(salon_a.tenant):
            Transaction.objects.create(
                tenant=salon_a.tenant,
                kind=Transaction.Kind.INCOME,
                category=Transaction.IncomeCategory.SERVICE,
                label="Box braids",
                amount=Decimal("20000"),
                occurred_on=date(2027, 1, 1),
            )

        login(api_client, salon_a.owner)
        response = api_client.get("/api/v1/transactions/summary/")

    assert response.status_code == 200, response.data
    # La borne haute est le jour du salon, pas celui du serveur.
    assert response.data["to"] == "2027-01-01"
    assert Decimal(response.data["totals"]["income"]) == Decimal("20000")
    assert response.data["totals"]["count"] == 1


@pytest.mark.django_db
def test_the_current_month_column_is_the_salons_month(api_client, salon_a):
    """Le mois courant est celui du salon.

    Le dernier jour du mois, un serveur en UTC et un salon a Shanghai ne sont
    pas dans le meme mois pendant huit heures : la colonne du mois qui vient
    de commencer manquerait, avec les recettes qu'elle porte.
    """
    with horloge_de_shanghai(salon_a):
        login(api_client, salon_a.owner)
        response = api_client.get("/api/v1/transactions/summary/")

    assert response.status_code == 200, response.data
    assert response.data["monthly"][-1]["month"] == "2027-01-01"


@pytest.mark.django_db
def test_an_entry_dated_today_is_accepted_even_when_utc_lags(api_client, salon_a):
    """Saisir une depense du jour ne doit pas etre refuse comme « future ».

    La validation comparait a la date du serveur : pendant les huit heures ou
    UTC est encore la veille, un salon de Shanghai se voyait refuser une
    depense bien reelle, avec un message incomprehensible.
    """
    with horloge_de_shanghai(salon_a):
        login(api_client, salon_a.owner)
        response = api_client.post(
            "/api/v1/transactions/",
            entry(salon_a, occurred_on="2027-01-01"),
            format="json",
            headers={"X-Tenant-Id": str(salon_a.tenant.id)},
        )

    assert response.status_code == 201, response.data
